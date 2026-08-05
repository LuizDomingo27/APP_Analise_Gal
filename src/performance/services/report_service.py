# -*- coding: utf-8 -*-
"""
Relatórios exportáveis da análise de performance (PDF e Excel).

Seguem hex fixo, sem tokens `--ag-*`: relatório impresso é branco em qualquer
tema — mesma regra já adotada por `src/ui/preview.py` e `src/services/exporter.py`.
"""

from __future__ import annotations

import re
from io import BytesIO
from xml.sax.saxutils import escape

import pandas as pd

from src.performance.domain.exceptions import ReportGenerationError

# Caracteres que o Excel recusa (openpyxl levanta IllegalCharacterError).
_ILLEGAL_SHEET_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# Prefixos que fazem o Excel avaliar a célula como fórmula/comando DDE.
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")


def sheet_text(value: object) -> str:
    """Texto seguro para planilha: sem caractere de controle e nunca avaliado como fórmula."""
    text = _ILLEGAL_SHEET_CHARS.sub("", str(value))
    if text.startswith(_FORMULA_TRIGGERS):
        text = "'" + text
    return text


def pdf_text(value: object) -> str:
    """Texto seguro para Paragraph do ReportLab, cuja marcação é lida como XML."""
    return escape(_ILLEGAL_SHEET_CHARS.sub("", str(value)))


def _number_br(value: float, decimals: int = 0) -> str:
    text = f"{float(value):,.{decimals}f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")


def _money_br(value: float) -> str:
    return "R$ " + _number_br(value, 2)


class PdfReportService:
    """Relatório gerencial das ordens críticas, pronto para impressão."""

    @staticmethod
    def build(
        critical_orders: pd.DataFrame,
        supplier_summary: pd.DataFrame,
        target: float,
        loss_threshold: float,
        period_label: str,
    ) -> bytes:
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_LEFT
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
            )

            output = BytesIO()
            document = SimpleDocTemplate(
                output,
                pagesize=landscape(A4),
                leftMargin=12 * mm, rightMargin=12 * mm,
                topMargin=12 * mm, bottomMargin=12 * mm,
                title="Relatório de Performance de Fornecedores",
            )
            styles = getSampleStyleSheet()
            title = ParagraphStyle(
                "TitleCustom", parent=styles["Title"], fontName="Helvetica-Bold",
                fontSize=20, leading=24, textColor=colors.HexColor("#0D1B17"), alignment=TA_LEFT,
            )
            subtitle = ParagraphStyle(
                "SubtitleCustom", parent=styles["BodyText"], fontSize=9,
                leading=12, textColor=colors.HexColor("#4A5752"),
            )
            cell = ParagraphStyle("Cell", parent=styles["BodyText"], fontSize=7, leading=9)
            cell_bold = ParagraphStyle("CellBold", parent=cell, fontName="Helvetica-Bold")

            story = [
                Paragraph("PERFORMANCE DE FORNECEDORES", title),
                Paragraph(
                    f"Período: {pdf_text(period_label)} &nbsp;&nbsp;|&nbsp;&nbsp; Meta: {target:.1%} "
                    f"&nbsp;&nbsp;|&nbsp;&nbsp; Ordens com perda a partir de {loss_threshold:.1%}",
                    subtitle,
                ),
                Spacer(1, 6 * mm),
            ]

            below = supplier_summary[supplier_summary["performance"] < target]
            empty = critical_orders.empty
            kpis = [
                ["Fornecedores analisados", "Abaixo da meta", "Ordens críticas",
                 "Peças perdidas", "Minutos gerados", "Custo das perdas"],
                [
                    str(len(supplier_summary)), str(len(below)), str(len(critical_orders)),
                    "0" if empty else _number_br(critical_orders["defect_pieces"].sum()),
                    "0,0" if empty else _number_br(critical_orders["minutes_generated"].sum(), 1),
                    _money_br(0 if empty else critical_orders["defect_cost"].sum()),
                ],
            ]
            kpi_table = Table(kpis, colWidths=[44 * mm] * 6)
            kpi_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1B17")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F2F7F5")),
                ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#0D1B17")),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 7.5),
                ("FONTSIZE", (0, 1), (-1, 1), 13),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E8EFEC")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([
                kpi_table, Spacer(1, 6 * mm),
                Paragraph("ORDENS CRÍTICAS", cell_bold), Spacer(1, 2 * mm),
            ])

            headers = ["Fornecedor", "Ordem mestre", "Data", "Real cortado", "Perdas",
                       "% perda", "Performance", "Minutos", "Custo (R$)"]
            data = [[Paragraph(header, cell_bold) for header in headers]]
            for row in critical_orders.itertuples(index=False):
                data.append([
                    Paragraph(pdf_text(row.supplier), cell),
                    Paragraph(pdf_text(row.master_order), cell),
                    row.production_period,
                    _number_br(row.real_cut),
                    _number_br(row.defect_pieces),
                    f"{row.defect_rate:.1%}",
                    f"{row.performance:.1%}",
                    _number_br(row.minutes_generated, 1),
                    _number_br(row.defect_cost, 2),
                ])
            if len(data) == 1:
                data.append(
                    [Paragraph("Nenhuma ordem encontrada para o corte selecionado.", cell)]
                    + [""] * 8
                )

            table = Table(
                data, repeatRows=1,
                colWidths=[64 * mm, 25 * mm, 24 * mm, 25 * mm, 20 * mm,
                           19 * mm, 23 * mm, 22 * mm, 25 * mm],
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#00B884")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F7F5")]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E8EFEC")),
                ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(table)

            if not below.empty:
                story.extend([
                    PageBreak(),
                    Paragraph("FORNECEDORES ABAIXO DA META", title),
                    Spacer(1, 4 * mm),
                ])
                supplier_data = [[
                    "Fornecedor", "Ordens", "Real cortado", "Peças perdidas",
                    "% perda", "Performance", "Minutos", "Custo (R$)",
                ]]
                for row in below.itertuples(index=False):
                    supplier_data.append([
                        Paragraph(pdf_text(row.supplier), cell), str(row.orders),
                        _number_br(row.real_cut), _number_br(row.defect_pieces),
                        f"{row.defect_rate:.1%}", f"{row.performance:.1%}",
                        _number_br(row.minutes_generated, 1), _number_br(row.defect_cost, 2),
                    ])
                supplier_table = Table(
                    supplier_data, repeatRows=1,
                    colWidths=[78 * mm, 20 * mm, 30 * mm, 28 * mm,
                               24 * mm, 27 * mm, 26 * mm, 28 * mm],
                )
                supplier_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1B17")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF4F5")]),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E8EFEC")),
                    ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.append(supplier_table)

            document.build(story)
            return output.getvalue()
        except Exception as exc:
            raise ReportGenerationError(
                "Não foi possível gerar o PDF para os filtros atuais."
            ) from exc


class SupplierExportService:
    """Exporta o ranking de fornecedores já filtrado."""

    HEADERS = [
        "Fornecedor", "Ordens", "Real cortado", "Peças com defeito",
        "Peças aprovadas", "Minutos gerados", "Custo das perdas (R$)",
        "Índice de defeitos", "Performance", "Status",
    ]

    @staticmethod
    def build_excel(supplier_summary: pd.DataFrame, target: float, period_label: str) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.formatting.rule import CellIsRule
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.worksheet.table import Table as ExcelTable, TableStyleInfo

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Fornecedores"
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "A5"

            sheet.merge_cells("A1:J1")
            sheet["A1"] = "PERFORMANCE DE FORNECEDORES"
            sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
            sheet["A1"].fill = PatternFill("solid", fgColor="0D1B17")
            sheet["A1"].alignment = Alignment(vertical="center")
            sheet.row_dimensions[1].height = 34

            sheet.merge_cells("A2:J2")
            sheet["A2"] = (
                f"Período: {period_label}  |  Meta: {target:.1%}  |  "
                f"Fornecedores: {len(supplier_summary)}"
            )
            sheet["A2"].font = Font(name="Aptos", size=10, color="4A5752")
            sheet["A2"].fill = PatternFill("solid", fgColor="F2F7F5")
            sheet["A2"].alignment = Alignment(vertical="center")
            sheet.row_dimensions[2].height = 24

            header_row = 4
            for column, header in enumerate(SupplierExportService.HEADERS, start=1):
                cell = sheet.cell(header_row, column, header)
                cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="00B884")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[header_row].height = 30

            for row_index, row in enumerate(supplier_summary.itertuples(index=False), start=5):
                values = [
                    sheet_text(row.supplier), int(row.orders), float(row.real_cut),
                    float(row.defect_pieces), float(row.approved_pieces),
                    float(row.minutes_generated), float(row.defect_cost),
                    float(row.defect_rate), float(row.performance), sheet_text(row.status),
                ]
                for column, value in enumerate(values, start=1):
                    cell = sheet.cell(row_index, column, value)
                    cell.font = Font(name="Aptos", size=10, color="0D1B17")
                    cell.alignment = Alignment(vertical="center", wrap_text=column in {1, 10})
                for column in (2, 3, 4, 5):
                    sheet.cell(row_index, column).number_format = "#,##0"
                sheet.cell(row_index, 6).number_format = "#,##0.0"
                sheet.cell(row_index, 7).number_format = 'R$ #,##0.00'
                for column in (8, 9):
                    sheet.cell(row_index, column).number_format = "0.0%"

            last_row = max(5, 4 + len(supplier_summary))
            if supplier_summary.empty:
                sheet.cell(5, 1, "Nenhum fornecedor encontrado para os filtros atuais.")

            table = ExcelTable(displayName="RankingFornecedores", ref=f"A4:J{last_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False,
            )
            sheet.add_table(table)

            # Performance é a coluna I. O realce condicional acompanha a meta
            # escolhida na tela, então a planilha exportada conta a mesma história
            # que o ranking em que o usuário clicou.
            if not supplier_summary.empty:
                sheet.conditional_formatting.add(
                    f"I5:I{last_row}",
                    CellIsRule(
                        operator="lessThan", formula=[str(target)],
                        fill=PatternFill("solid", fgColor="FFF1F3"),
                        font=Font(color="B83F50", bold=True),
                    ),
                )
                sheet.conditional_formatting.add(
                    f"I5:I{last_row}",
                    CellIsRule(
                        operator="greaterThanOrEqual", formula=[str(target)],
                        fill=PatternFill("solid", fgColor="EAF9F5"),
                        font=Font(color="00805C", bold=True),
                    ),
                )

            widths = {"A": 46, "B": 12, "C": 18, "D": 20, "E": 18,
                      "F": 18, "G": 22, "H": 19, "I": 16, "J": 20}
            for column, width in widths.items():
                sheet.column_dimensions[column].width = width

            separator = Side(style="thin", color="E8EFEC")
            for row in sheet.iter_rows(min_row=5, max_row=last_row, min_col=1, max_col=10):
                for cell in row:
                    cell.border = Border(bottom=separator)

            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        except Exception as exc:
            raise ReportGenerationError(
                "Não foi possível gerar o Excel de fornecedores."
            ) from exc

    @staticmethod
    def build_pdf(supplier_summary: pd.DataFrame, target: float, period_label: str) -> bytes:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
            )

            output = BytesIO()
            document = SimpleDocTemplate(
                output,
                pagesize=landscape(A4),
                leftMargin=12 * mm, rightMargin=12 * mm,
                topMargin=12 * mm, bottomMargin=14 * mm,
                title="Ranking de Performance de Fornecedores",
            )
            styles = getSampleStyleSheet()
            title = ParagraphStyle(
                "SupplierTitle", parent=styles["Title"], fontName="Helvetica-Bold",
                fontSize=18, leading=22, textColor=colors.HexColor("#0D1B17"), alignment=0,
            )
            subtitle = ParagraphStyle(
                "SupplierSubtitle", parent=styles["BodyText"], fontSize=9,
                leading=12, textColor=colors.HexColor("#4A5752"),
            )
            cell = ParagraphStyle("SupplierCell", parent=styles["BodyText"], fontSize=7, leading=9)
            cell_bold = ParagraphStyle("SupplierCellBold", parent=cell, fontName="Helvetica-Bold")

            below_target = (
                int(supplier_summary["performance"].lt(target).sum())
                if not supplier_summary.empty else 0
            )
            total_cost = (
                float(supplier_summary["defect_cost"].sum())
                if not supplier_summary.empty else 0.0
            )
            story = [
                Paragraph("RANKING DE PERFORMANCE DE FORNECEDORES", title),
                Paragraph(
                    f"Período: {pdf_text(period_label)} &nbsp;&nbsp;|&nbsp;&nbsp; Meta: {target:.1%}",
                    subtitle,
                ),
                Spacer(1, 4 * mm),
            ]

            kpis = [
                ["Fornecedores analisados", "Abaixo da meta", "Dentro da meta", "Custo das perdas"],
                [
                    str(len(supplier_summary)), str(below_target),
                    str(len(supplier_summary) - below_target), _money_br(total_cost),
                ],
            ]
            kpi_table = Table(kpis, colWidths=[49 * mm] * 4)
            kpi_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1B17")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F2F7F5")),
                ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#0D1B17")),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("FONTSIZE", (0, 1), (-1, 1), 13),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E8EFEC")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([kpi_table, Spacer(1, 5 * mm)])

            headers = ["Fornecedor", "Ordens", "Real cortado", "Defeitos", "Aprovadas",
                       "Minutos", "Custo (R$)", "% defeito", "Performance", "Status"]
            data = [[Paragraph(header, cell_bold) for header in headers]]
            for row in supplier_summary.itertuples(index=False):
                data.append([
                    Paragraph(pdf_text(row.supplier), cell),
                    _number_br(row.orders), _number_br(row.real_cut),
                    _number_br(row.defect_pieces), _number_br(row.approved_pieces),
                    _number_br(row.minutes_generated, 1), _number_br(row.defect_cost, 2),
                    f"{row.defect_rate:.1%}", f"{row.performance:.1%}",
                    Paragraph(pdf_text(row.status), cell),
                ])
            if len(data) == 1:
                data.append(
                    [Paragraph("Nenhum fornecedor encontrado para os filtros atuais.", cell)]
                    + [""] * 9
                )

            supplier_table = Table(
                data, repeatRows=1,
                colWidths=[57 * mm, 16 * mm, 25 * mm, 21 * mm, 22 * mm,
                           21 * mm, 26 * mm, 20 * mm, 24 * mm, 26 * mm],
            )
            supplier_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#00B884")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFCFB")]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E8EFEC")),
                ("ALIGN", (1, 1), (8, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(supplier_table)

            def add_page_number(canvas, doc) -> None:
                canvas.saveState()
                canvas.setFillColor(colors.HexColor("#7C8985"))
                canvas.setFont("Helvetica", 8)
                canvas.drawRightString(285 * mm, 7 * mm, f"Página {doc.page}")
                canvas.restoreState()

            document.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)
            return output.getvalue()
        except Exception as exc:
            raise ReportGenerationError(
                "Não foi possível gerar o PDF de fornecedores."
            ) from exc
