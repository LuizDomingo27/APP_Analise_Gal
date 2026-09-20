"""Resumo por lançamento e exportações, sem acesso ao banco."""
from io import BytesIO
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from src.config.settings import COLS


SUMMARY_COLUMNS = [
    "Nome fornecedor", "Data cobrança", "Data vencimento", "Data pagamento",
    "Total defeitos", "Total minutos", "Valor da cobrança", "Status",
]


def build_charge_summary(frames: list[pd.DataFrame]) -> pd.DataFrame:
    """Soma itens por código dentro de cada origem; não une cobranças do fornecedor."""
    rows = []
    for frame in frames:
        if frame is None or frame.empty:
            continue
        for _, group in frame.groupby("COD_LANCAMENTO", sort=False, dropna=False):
            first = group.iloc[0]
            dates = [pd.to_datetime(first.get(c), dayfirst=True, errors="coerce")
                     for c in ("DATA_COBRANCA", "DATA_VENCIMENTO", "DATA_PAGAMENTO")]
            sums = [pd.to_numeric(group[c], errors="coerce").sum()
                    for c in (COLS["quantity"], COLS["minutes"], COLS["value_brl"])]
            rows.append([first[COLS["supplier"]], *dates, *sums, first.get(COLS["status"], "Pendente")])
    return pd.DataFrame(rows, columns=SUMMARY_COLUMNS).sort_values(
        ["Nome fornecedor", "Data cobrança"], ascending=[True, False], ignore_index=True,
    )


def summary_excel(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo de cobranças"
    ws.append(SUMMARY_COLUMNS)
    for row in df[SUMMARY_COLUMNS].itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v for v in row])
        # Nomes são texto literal, inclusive quando começam por '='.
        ws.cell(ws.max_row, 1).data_type = "s"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="00805C")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30
    for column, width in zip("ABCDEFGH", [42, 19, 19, 19, 18, 19, 23, 16]):
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row[0].row].height = 32
        for cell in row[1:4]:
            cell.number_format = "dd/mm/yyyy"
        row[4].number_format = "#,##0"
        row[5].number_format = "#,##0.00"
        row[6].number_format = '"R$" #,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def summary_pdf(df: pd.DataFrame) -> bytes:
    out = BytesIO()
    styles = getSampleStyleSheet()
    normal = styles["BodyText"]
    normal.fontSize = 8
    normal.leading = 11
    header = normal.clone("SummaryHeader", textColor=colors.white)
    data = [[Paragraph(escape(c), header) for c in SUMMARY_COLUMNS]]
    for row in df[SUMMARY_COLUMNS].itertuples(index=False, name=None):
        dates = ["" if pd.isna(v) else pd.Timestamp(v).strftime("%d/%m/%Y") for v in row[1:4]]
        numbers = [f"{v:,.{precision}f}".replace(",", "_").replace(".", ",").replace("_", ".")
                   for v, precision in zip(row[4:7], [0, 2, 2])]
        numbers[-1] = "R$ " + numbers[-1]
        data.append([Paragraph(escape(str(row[0])), normal), *dates, *numbers,
                     Paragraph(escape(str(row[7])), normal)])
    table = Table(data, colWidths=[172, 78, 78, 78, 68, 70, 90, 76], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#00805C")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF7F3")]),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (4, 1), (6, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(810, 18, f"Página {doc.page}")
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=30,
                            rightMargin=30, topMargin=30, bottomMargin=30)
    doc.build([Paragraph("Resumo de cobranças por fornecedor", styles["Title"]),
               Paragraph(f"{len(df)} cobrança(s) no recorte selecionado", normal),
               Spacer(1, 14), table], onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
