# -*- coding: utf-8 -*-
"""Regressões da auditoria de cobertura da página de Performance."""
from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.performance.presentation.formatters import (
    currency_br,
    date_br,
    number_br,
    percent_br,
    to_float,
)
from src.performance.services.report_service import (
    PdfReportService,
    SupplierExportService,
    pdf_text,
    sheet_text,
)


def summary_row(**overrides) -> pd.DataFrame:
    row = {
        "supplier": "FORNECEDOR X", "orders": 1, "real_cut": 100.0,
        "defect_pieces": 5.0, "approved_pieces": 95.0, "minutes_generated": 10.0,
        "defect_cost": 12.5, "defect_rate": 0.05, "performance": 0.95,
        "status": "Dentro da meta",
    }
    row.update(overrides)
    return pd.DataFrame([row])


# --- formatadores: nenhum float() sem guarda --------------------------------

@pytest.mark.parametrize("value", [None, pd.NA, pd.NaT, "abc", float("nan"), float("inf")])
def test_percent_br_nunca_estoura(value):
    assert percent_br(value) == "—"


@pytest.mark.parametrize("value", [None, pd.NA, "abc", float("inf")])
def test_currency_br_nunca_estoura(value):
    assert currency_br(value) == "—"


@pytest.mark.parametrize("value", [None, pd.NA, pd.NaT, "abc", object()])
def test_number_br_nunca_estoura(value):
    assert number_br(value) == "—"


def test_date_br_trata_nat():
    assert date_br(pd.NaT) == "—"


def test_to_float_rejeita_nao_finito():
    assert to_float(float("inf")) is None
    assert to_float("2.5") == 2.5


# --- Excel: injeção de fórmula e caracteres ilegais -------------------------

@pytest.mark.parametrize("payload", [
    '=HYPERLINK("http://evil","clique")',
    "=cmd|'/c calc'!A1",
    "+1+1",
    "-2+3",
    "@SUM(A1)",
])
def test_export_excel_nunca_grava_formula_viva(payload):
    data = SupplierExportService.build_excel(summary_row(supplier=payload), 0.9, "01/01/2026")
    cell = load_workbook(io.BytesIO(data)).active.cell(5, 1)
    assert cell.data_type != "f", f"{payload!r} virou fórmula executável"
    assert str(cell.value).startswith("'")


def test_sheet_text_remove_caracteres_de_controle():
    assert sheet_text("AC\x01ME\x00") == "ACME"


def test_export_excel_sobrevive_a_caractere_de_controle():
    data = SupplierExportService.build_excel(summary_row(supplier="AC\x01ME"), 0.9, "01/01/2026")
    assert load_workbook(io.BytesIO(data)).active.cell(5, 1).value == "ACME"


def test_pdf_text_escapa_marcacao_e_remove_controle():
    assert pdf_text("<b>&x</b>\x01") == "&lt;b&gt;&amp;x&lt;/b&gt;"


def test_pdf_sobrevive_a_nome_hostil():
    hostil = "</script><img src=x onerror=alert(1)>&\x01"
    summary = summary_row(supplier=hostil)
    assert SupplierExportService.build_pdf(summary, 0.9, hostil).startswith(b"%PDF")


# --- shim de renderização ---------------------------------------------------

def test_render_html_usa_st_iframe_quando_existe(monkeypatch):
    import streamlit as st

    from src.performance.presentation import components as ui

    calls = []
    monkeypatch.setattr(st, "iframe", lambda m, **kw: calls.append((m, kw)), raising=False)
    ui.render_html("<p>x</p>", 100, scrolling=True)
    assert calls == [("<p>x</p>", {"height": 100})]


def test_render_html_cai_para_components_html(monkeypatch):
    import streamlit as st
    import streamlit.components.v1 as components

    from src.performance.presentation import components as ui

    calls = []
    monkeypatch.delattr(st, "iframe", raising=False)
    monkeypatch.setattr(components, "html", lambda m, **kw: calls.append((m, kw)))
    ui.render_html("<p>x</p>", 100, scrolling=True)
    assert calls == [("<p>x</p>", {"height": 100, "scrolling": True})]
