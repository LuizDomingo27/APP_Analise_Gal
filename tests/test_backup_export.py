# -*- coding: utf-8 -*-
"""
Testes do backup completo do banco (src/data/backup.py, src/services/
backup_exporter.py e src/ui/backup.py).

O backup é a ÚNICA cópia dos dados fora do Supabase enquanto o plano for o
gratuito, então os testes cobrem três riscos, nesta ordem de gravidade:

  1. Backup silenciosamente inútil — arquivo gerado, mas com data virando
     texto, número virando string ou célula truncada sem aviso. Daí os testes
     de round-trip de tipos e de number_format.
  2. Backup incompleto sem ninguém perceber — uma tabela falha e leva o resto
     junto, ou some do pacote sem deixar rastro. Daí o isolamento por tabela e
     o manifesto _INFO.
  3. Vazamento — a tabela de usuários (com hash de senha) só pode ser lida por
     administrador. Daí o teste do gate.
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy import text

import src.auth.auth_db as auth_db
import src.data.backup as bk
import src.data.database as db
import src.services.backup_exporter as bx
import src.ui.backup as bkui
from src.config.settings import COLS


# ── Fixtures ──────────────────────────────────────────────────────────────────

@pytest.fixture
def temp_db(tmp_path, monkeypatch):
    """Isola o backup num SQLite temporário com o schema real do app."""
    sqlite_url = f"sqlite:///{tmp_path / 'test_backup.db'}"
    monkeypatch.setenv("DATABASE_URL", sqlite_url)
    monkeypatch.setattr(db, "_database_url", lambda: sqlite_url)
    db.get_engine.clear()
    db.create_tables.clear()
    auth_db.create_users_table.clear()
    auth_db.list_users.clear()
    db.create_tables()
    auth_db.create_users_table()
    yield
    db.get_engine.clear()
    db.create_tables.clear()
    auth_db.create_users_table.clear()
    auth_db.list_users.clear()


def _registro(data="2026-07-01", fornecedor="Oficina A", qtd=10, valor=150.5):
    return {
        COLS["date"]: data,
        COLS["order"]: "OM-100",
        COLS["material"]: "TECIDO X",
        COLS["supplier"]: fornecedor,
        COLS["quantity"]: qtd,
        COLS["location"]: "MANGA",
        COLS["defect"]: "PONTO ESTOURADO",
        COLS["real_cut"]: "5",
        COLS["pct_remonte"]: 0.125,
        COLS["key"]: "CHV-1",
        COLS["process_time"]: "00:10",
        COLS["minutes"]: 12.5,
        COLS["value_brl"]: valor,
        COLS["status"]: "Pendente",
    }


def _inserir(tabela: str, linhas: list[dict]) -> None:
    with db.get_connection() as conn:
        pd.DataFrame(linhas).to_sql(tabela, conn, if_exists="append", index=False)
        conn.commit()


def _resultado(nome="registros_defeitos", df=None, status=None, colunas=None, erro=""):
    """ResultadoTabela pronto, sem precisar de banco."""
    df = pd.DataFrame([_registro()]) if df is None else df
    if status is None:
        status = bk.STATUS_OK if len(df) else bk.STATUS_VAZIA
    return bk.ResultadoTabela(
        nome=nome,
        rotulo="Rótulo de teste",
        status=status,
        df=None if status in (bk.STATUS_AUSENTE, bk.STATUS_ERRO) else df,
        colunas=colunas or [(c, "TEXT") for c in df.columns],
        erro=erro,
    )


def _abrir(dados: bytes):
    return load_workbook(io.BytesIO(dados))


def _linha_cabecalho(ws) -> int:
    """3 normalmente; 4 quando a planilha tem a faixa de aviso (UserGal)."""
    primeira = ws.cell(row=3, column=1).value
    return 4 if isinstance(primeira, str) and primeira.startswith("⚠") else 3


def _ler_como_df(dados: bytes) -> pd.DataFrame:
    ws = _abrir(dados).worksheets[0]
    return pd.read_excel(io.BytesIO(dados), header=_linha_cabecalho(ws) - 1)


# ── Coleta: leitura das tabelas ───────────────────────────────────────────────

def test_coleta_le_todas_as_tabelas_do_schema(temp_db):
    _inserir("registros_defeitos", [_registro(), _registro(data="2026-07-02")])

    resultados = bk.coletar_backup()
    por_nome = {r.nome: r for r in resultados}

    assert [t.nome for t in bk.TABELAS_BACKUP] == [r.nome for r in resultados]
    assert por_nome["registros_defeitos"].status == bk.STATUS_OK
    assert por_nome["registros_defeitos"].linhas == 2
    # Tabelas criadas pelo schema mas ainda sem dados existem e são exportadas.
    assert por_nome["devolucoes"].status == bk.STATUS_VAZIA
    assert por_nome["devolucoes"].exportavel


def test_coleta_traz_a_coluna_id_para_restauracao(temp_db):
    """O id é chave da linha no destino — um backup sem ele não restaura igual."""
    _inserir("registros_defeitos", [_registro()])
    resultado = bk.carregar_tabela(bk.TabelaBackup("registros_defeitos", "x"))
    assert "id" in resultado.df.columns


def test_coleta_registra_tipos_das_colunas_para_o_manifesto(temp_db):
    resultado = bk.carregar_tabela(bk.TabelaBackup("registros_defeitos", "x"))
    tipos = dict(resultado.colunas)
    assert tipos[COLS["quantity"]].upper().startswith("INTEGER")
    assert "TEXT" in tipos[COLS["supplier"]].upper()


def test_tabela_ausente_nao_derruba_a_coleta(temp_db):
    with db.get_connection() as conn:
        conn.execute(text("DROP TABLE devolucoes"))
        conn.commit()

    resultados = bk.coletar_backup()
    por_nome = {r.nome: r for r in resultados}

    assert por_nome["devolucoes"].status == bk.STATUS_AUSENTE
    assert not por_nome["devolucoes"].exportavel
    assert por_nome["registros_defeitos"].exportavel  # as demais seguem normais


def test_falha_em_uma_tabela_nao_impede_as_outras(temp_db, monkeypatch):
    original = pd.read_sql

    def _read_sql_quebrado(sql, con, *args, **kwargs):
        if "historico_cobrancas" in str(sql):
            raise RuntimeError("timeout na leitura")
        return original(sql, con, *args, **kwargs)

    monkeypatch.setattr(bk.pd, "read_sql", _read_sql_quebrado)

    resultados = bk.coletar_backup()
    por_nome = {r.nome: r for r in resultados}

    assert por_nome["historico_cobrancas"].status == bk.STATUS_ERRO
    assert por_nome["historico_cobrancas"].erro
    assert sum(1 for r in resultados if r.exportavel) == len(bk.TABELAS_BACKUP) - 1


def test_callback_de_progresso_que_falha_nao_cancela_o_backup(temp_db):
    def _progresso_ruim(*_a):
        raise ValueError("widget morreu no meio")

    resultados = bk.coletar_backup(progresso=_progresso_ruim)
    assert len(resultados) == len(bk.TABELAS_BACKUP)


def test_usuarios_entram_no_backup_com_hash_e_salt(temp_db):
    ok, _ = auth_db.create_user("admin.teste", "Admin Teste", "senha123", role="admin")
    assert ok

    resultado = bk.carregar_tabela(bk.TabelaBackup("UserGal", "Usuários"))

    assert resultado.status == bk.STATUS_OK
    assert {"senha_hash", "salt", "role"} <= set(resultado.df.columns)
    assert resultado.df.loc[0, "senha_hash"]


# ── Pacote .zip ───────────────────────────────────────────────────────────────

def test_zip_tem_manifesto_e_um_arquivo_por_tabela(temp_db):
    _inserir("registros_defeitos", [_registro()])
    resultados = bk.coletar_backup()

    with zipfile.ZipFile(io.BytesIO(bx.build_backup_zip(resultados))) as zf:
        nomes = set(zf.namelist())

    assert "_INFO.xlsx" in nomes
    for r in resultados:
        assert (f"{r.nome}.xlsx" in nomes) is r.exportavel


def test_tabela_com_erro_fica_fora_do_zip_mas_consta_no_manifesto():
    resultados = [
        _resultado("registros_defeitos"),
        _resultado("devolucoes", status=bk.STATUS_ERRO, erro="timeout na leitura"),
    ]

    with zipfile.ZipFile(io.BytesIO(bx.build_backup_zip(resultados))) as zf:
        assert "devolucoes.xlsx" not in zf.namelist()
        info = _ler_como_df(zf.read("_INFO.xlsx"))

    # Cabeçalho sai em maiúsculas, como nos demais exports do app.
    linha = info[info["TABELA"] == "devolucoes"].iloc[0]
    assert linha["STATUS"] == "Falhou"
    assert "timeout" in linha["OBSERVAÇÃO"]


def test_imagens_ficam_fora_do_zip_e_o_manifesto_explica_o_motivo():
    """Ausência por decisão do time não pode parecer falha na restauração."""
    with zipfile.ZipFile(io.BytesIO(bx.build_backup_zip([_resultado()]))) as zf:
        assert "defeitos_imagens.xlsx" not in zf.namelist()
        info = _ler_como_df(zf.read("_INFO.xlsx"))

    linha = info[info["TABELA"] == "defeitos_imagens"].iloc[0]
    assert linha["STATUS"] == "Não incluída (decisão)"
    assert "recadastradas" in linha["OBSERVAÇÃO"]


def test_manifesto_lista_os_tipos_de_cada_coluna():
    resultados = [_resultado("registros_defeitos", colunas=[("id", "BIGINT"), ("FORNECEDOR", "TEXT")])]

    with zipfile.ZipFile(io.BytesIO(bx.build_backup_zip(resultados))) as zf:
        info_bytes = zf.read("_INFO.xlsx")

    ws = _abrir(info_bytes)["Colunas"]
    colunas = pd.read_excel(io.BytesIO(info_bytes), sheet_name="Colunas",
                            header=_linha_cabecalho(ws) - 1)
    assert list(colunas["COLUNA"]) == ["id", "FORNECEDOR"]
    assert list(colunas["TIPO NO BANCO"]) == ["BIGINT", "TEXT"]


def test_manifesto_ensina_a_reimportar_sem_perder_zeros_a_esquerda():
    """'00123' relido com inferência automática vira 123 e ninguém percebe."""
    with zipfile.ZipFile(io.BytesIO(bx.build_backup_zip([_resultado()]))) as zf:
        info_bytes = zf.read("_INFO.xlsx")

    wb = _abrir(info_bytes)
    assert "Como restaurar" in wb.sheetnames
    texto = " ".join(
        str(c.value) for linha in wb["Como restaurar"].iter_rows() for c in linha
    )
    assert "dtype=str" in texto
    assert "ORDEM MESTRE" in texto


def test_nome_do_arquivo_tem_data_e_hora():
    nome = bx.nome_do_arquivo(datetime(2026, 7, 24, 9, 5))
    assert nome == "backup_analise_gal_20260724_0905.zip"


# ── Tipos: o ponto que torna o backup restaurável ─────────────────────────────

def test_round_trip_preserva_os_tipos_de_dado(temp_db):
    _inserir("registros_defeitos", [_registro(), _registro(data="2026-07-02", qtd=3)])
    resultado = bk.carregar_tabela(bk.TabelaBackup("registros_defeitos", "x"))

    df = _ler_como_df(bx.build_table_xlsx(resultado, datetime.now()))

    assert pd.api.types.is_datetime64_any_dtype(df[COLS["date"]])
    assert pd.api.types.is_integer_dtype(df[COLS["quantity"]])
    assert pd.api.types.is_float_dtype(df[COLS["value_brl"]])
    assert df.loc[0, COLS["date"]] == pd.Timestamp("2026-07-01")
    assert df.loc[0, COLS["value_brl"]] == pytest.approx(150.5)


def test_formatos_de_numero_aplicados_nas_colunas_certas(temp_db):
    _inserir("registros_defeitos", [_registro()])
    resultado = bk.carregar_tabela(bk.TabelaBackup("registros_defeitos", "x"))

    ws = _abrir(bx.build_table_xlsx(resultado, datetime.now())).worksheets[0]
    cab = _linha_cabecalho(ws)
    col = {ws.cell(row=cab, column=c).value: c for c in range(1, ws.max_column + 1)}
    fmt = lambda nome: ws.cell(row=cab + 1, column=col[nome.upper()]).number_format

    assert fmt(COLS["value_brl"]) == bx.FMT_MONEY
    assert fmt(COLS["minutes"]) == bx.FMT_DECIMAL
    assert fmt(COLS["quantity"]) == bx.FMT_INT
    assert fmt(COLS["pct_remonte"]) == bx.FMT_PERCENT
    assert fmt(COLS["date"]) == bx.FMT_DATE


def test_percentual_vai_como_fracao_para_o_formato_ler(temp_db):
    """0,125 com formato 0.00% aparece como 12,50% — multiplicar aqui dobraria."""
    _inserir("registros_defeitos", [_registro()])
    resultado = bk.carregar_tabela(bk.TabelaBackup("registros_defeitos", "x"))

    ws = _abrir(bx.build_table_xlsx(resultado, datetime.now())).worksheets[0]
    cab = _linha_cabecalho(ws)
    col = [c for c in range(1, ws.max_column + 1)
           if ws.cell(row=cab, column=c).value == COLS["pct_remonte"].upper()][0]
    assert ws.cell(row=cab + 1, column=col).value == pytest.approx(0.125)


def test_identificador_textual_nao_vira_numero():
    """'00123' em ORDEM MESTRE não pode perder os zeros à esquerda."""
    df = pd.DataFrame({COLS["order"]: ["00123", "00456"]})
    ws = _abrir(bx.build_table_xlsx(_resultado(df=df), datetime.now())).worksheets[0]
    assert ws.cell(row=4, column=1).value == "00123"


def test_coluna_de_data_com_valor_sujo_fica_como_texto():
    """Conversão é tudo-ou-nada: melhor texto fiel do que data adivinhada."""
    df = pd.DataFrame({COLS["date"]: ["2026-07-01", "sem data"]})
    ws = _abrir(bx.build_table_xlsx(_resultado(df=df), datetime.now())).worksheets[0]
    assert ws.cell(row=4, column=1).value == "2026-07-01"
    assert ws.cell(row=5, column=1).value == "sem data"


def test_valor_ausente_vira_celula_vazia_e_nao_o_texto_nan():
    df = pd.DataFrame({COLS["supplier"]: ["Oficina A", None], COLS["value_brl"]: [10.0, None]})
    ws = _abrir(bx.build_table_xlsx(_resultado(df=df), datetime.now())).worksheets[0]
    assert ws.cell(row=5, column=1).value is None
    assert ws.cell(row=5, column=2).value is None


def test_data_vazia_no_banco_continua_exportavel():
    """DATA_PAGAMENTO só é preenchida na baixa — a coluna toda vazia é normal."""
    df = pd.DataFrame({"DATA_PAGAMENTO": [None, ""]})
    ws = _abrir(bx.build_table_xlsx(_resultado(df=df), datetime.now())).worksheets[0]
    assert ws.cell(row=4, column=1).value is None


# ── Limites do formato xlsx (arquivo corrompido é backup perdido) ─────────────

def test_texto_gigante_e_truncado_em_vez_de_corromper_o_arquivo():
    df = pd.DataFrame({COLS["supplier"]: ["x" * 40_000]})
    ws = _abrir(bx.build_table_xlsx(_resultado(df=df), datetime.now())).worksheets[0]
    valor = ws.cell(row=4, column=1).value
    assert len(valor) <= bx._MAX_CELL_CHARS
    assert valor.endswith(bx._TRUNC_MARK)


def test_caractere_de_controle_e_removido():
    """openpyxl levanta IllegalCharacterError e derrubaria o backup inteiro."""
    df = pd.DataFrame({COLS["supplier"]: ["Oficina\x0bA"]})
    ws = _abrir(bx.build_table_xlsx(_resultado(df=df), datetime.now())).worksheets[0]
    assert ws.cell(row=4, column=1).value == "OficinaA"


def test_binario_vira_descricao_legivel():
    assert "bytes" in bx.valor_celula(b"\x89PNG\r\n")


@pytest.mark.parametrize("nome, esperado", [
    ("registros_defeitos", "registros_defeitos"),
    ("tabela/com:caracteres*proibidos", "tabela-com-caracteres-proibidos"),
    ("n" * 40, "n" * 31),
])
def test_nome_de_aba_sempre_valido(nome, esperado):
    assert bx.nome_de_aba(nome) == esperado


def test_tabela_vazia_gera_planilha_so_com_cabecalho(temp_db):
    resultado = bk.carregar_tabela(bk.TabelaBackup("devolucoes", "Devoluções"))
    assert resultado.status == bk.STATUS_VAZIA

    ws = _abrir(bx.build_table_xlsx(resultado, datetime.now())).worksheets[0]
    cab = _linha_cabecalho(ws)
    assert ws.cell(row=cab, column=1).value == "ID"
    assert ws.cell(row=cab + 1, column=1).value is None


def test_planilha_de_usuarios_avisa_sobre_conteudo_sensivel():
    df = pd.DataFrame({"username": ["admin"], "senha_hash": ["abc"], "salt": ["def"]})
    ws = _abrir(bx.build_table_xlsx(_resultado("UserGal", df=df), datetime.now())).worksheets[0]
    aviso = ws.cell(row=3, column=1).value
    assert aviso.startswith("⚠") and "hash" in aviso
    assert ws.cell(row=4, column=1).value == "USERNAME"  # cabeçalho desceu uma linha


def test_planilha_tem_congelamento_e_autofiltro(temp_db):
    _inserir("registros_defeitos", [_registro()])
    resultado = bk.carregar_tabela(bk.TabelaBackup("registros_defeitos", "x"))
    ws = _abrir(bx.build_table_xlsx(resultado, datetime.now())).worksheets[0]
    assert ws.freeze_panes == "A4"
    assert ws.auto_filter.ref.startswith("A3:")


# ── UI: gate de administrador e fronteiras de erro ────────────────────────────

class _FakeProgress:
    def __init__(self):
        self.textos: list[str] = []

    def progress(self, _valor, text=""):
        self.textos.append(text)

    def empty(self):
        pass


class _FakeSt:
    """Recorder mínimo dos widgets usados por _gerar_backup/render_backup_section."""

    def __init__(self):
        self.errors: list[str] = []
        self.warnings: list[str] = []
        self.session_state: dict = {}
        self.barra = _FakeProgress()

    def progress(self, _valor, text=""):
        return self.barra

    def error(self, texto):
        self.errors.append(texto)

    def warning(self, texto):
        self.warnings.append(texto)

    def markdown(self, *_a, **_k):
        pass


class _FakeSession:
    def __init__(self, admin: bool):
        self._admin = admin

    def is_admin(self) -> bool:
        return self._admin

    def current_user(self) -> dict:
        return {"username": "admin.teste", "role": "admin" if self._admin else "user"}


@pytest.fixture
def ui(monkeypatch):
    fake = _FakeSt()
    monkeypatch.setattr(bkui, "st", fake)
    return fake


def test_usuario_comum_nao_ve_a_secao_nem_le_o_banco(ui, monkeypatch):
    """A seção lê o banco INTEIRO, incluindo os hashes: o gate é o controle."""
    monkeypatch.setattr(bkui, "session", _FakeSession(admin=False))
    chamadas = []
    monkeypatch.setattr(bkui, "coletar_backup", lambda **_k: chamadas.append(1))

    bkui.render_backup_section()  # não levanta

    assert chamadas == []
    assert ui.session_state == {}


def test_banco_indisponivel_mostra_mensagem_amigavel(ui, monkeypatch):
    def _explode(**_k):
        raise db.DatabaseUnavailableError("Não foi possível conectar ao banco.")

    monkeypatch.setattr(bkui, "coletar_backup", _explode)

    assert bkui._gerar_backup("admin.teste") is None
    assert "conectar ao banco" in ui.errors[0]
    assert not ui.warnings  # erro de banco não vira aviso genérico


def test_falha_ao_montar_o_zip_nao_derruba_a_pagina(ui, monkeypatch):
    monkeypatch.setattr(bkui, "coletar_backup", lambda **_k: [_resultado()])
    monkeypatch.setattr(
        bkui, "build_backup_zip", lambda *_a, **_k: (_ for _ in ()).throw(MemoryError())
    )

    assert bkui._gerar_backup("admin.teste") is None
    assert "Não foi possível gerar o backup" in ui.warnings[0]


def test_nenhuma_tabela_legivel_avisa_em_vez_de_baixar_zip_vazio(ui, monkeypatch):
    resultados = [_resultado("devolucoes", status=bk.STATUS_ERRO, erro="sem conexão")]
    monkeypatch.setattr(bkui, "coletar_backup", lambda **_k: resultados)

    assert bkui._gerar_backup("admin.teste") is None
    assert "Nenhuma tabela" in ui.errors[0]
    assert "sem conexão" in ui.errors[0]


def test_backup_bem_sucedido_devolve_resumo_para_a_sessao(ui, monkeypatch):
    resultados = [
        _resultado("registros_defeitos"),
        _resultado("devolucoes", status=bk.STATUS_ERRO, erro="timeout"),
    ]
    monkeypatch.setattr(bkui, "coletar_backup", lambda **_k: resultados)

    payload = bkui._gerar_backup("admin.teste")

    assert payload["arquivo"].startswith("backup_analise_gal_")
    assert payload["tabelas"] == 1
    assert payload["linhas"] == 1
    assert payload["falhas"] == ["devolucoes"]
    assert zipfile.is_zipfile(io.BytesIO(payload["bytes"]))
    assert ui.barra.textos  # houve feedback de progresso
