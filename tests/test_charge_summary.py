from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from src.config.settings import COLS
from src.data.cobranca_history import HISTORY_LABELS
from src.services.charge_summary_exporter import (
    SUMMARY_COLUMNS, build_charge_summary, summary_excel, summary_pdf,
)


def records():
    return pd.DataFrame({
        "COD_LANCAMENTO": ["A", "A", "B"],
        COLS["supplier"]: ["Oficina São José"] * 3,
        "DATA_COBRANCA": ["20/09/2026"] * 3,
        "DATA_VENCIMENTO": ["30/09/2026"] * 3,
        "DATA_PAGAMENTO": ["", "", "28/09/2026"],
        COLS["quantity"]: [2, 3, 4],
        COLS["minutes"]: [1.25, 2.5, 6],
        COLS["value_brl"]: [100.1, 200.2, 400],
        COLS["status"]: ["Pendente", "Pendente", "Pago"],
    })


def test_summary_preserves_distinct_charges_and_sums_items():
    result = build_charge_summary([records()])
    assert list(result.columns) == SUMMARY_COLUMNS
    assert len(result) == 2
    assert result["Total defeitos"].tolist() == [5, 4]
    assert result["Total minutos"].tolist() == [3.75, 6]
    assert round(result["Valor da cobrança"].sum(), 2) == 700.3
    assert result["Status"].tolist() == ["Pendente", "Pago"]
    assert pd.isna(result.iloc[0]["Data pagamento"])


def test_excel_keeps_numeric_dates_status_and_empty_payment():
    df = records()
    df[COLS["supplier"]] = "=Oficina"
    ws = load_workbook(BytesIO(summary_excel(build_charge_summary([df])))).active
    assert ws.max_row == 3
    assert ws.max_column == 8
    assert ws["A2"].data_type == "s"
    assert ws["B2"].value.day == 20
    assert ws["D2"].value is None
    assert ws["E2"].value == 5
    assert ws["H3"].value == "Pago"


def test_empty_summary_and_multipage_pdf():
    assert build_charge_summary([None, pd.DataFrame()]).empty
    result = build_charge_summary([records()])
    result = pd.concat([result] * 60, ignore_index=True)
    result["Nome fornecedor"] = "Oficina São José & Filhos com nome longo para quebra de linha"
    pdf = summary_pdf(result)
    assert pdf.startswith(b"%PDF")
    assert pdf.count(b"/Type /Page\n") > 1


def _summary_page():
    import streamlit as st
    from src.ui.charge_summary import render_summary_exports
    cols = st.columns(2)
    render_summary_exports(st.session_state["frame"], *cols, "pagamentos")


def test_summary_ui_has_two_downloads_and_paid_status():
    app = AppTest.from_function(_summary_page, default_timeout=60)
    app.session_state["frame"] = records().rename(columns=HISTORY_LABELS)
    app.run()
    assert not app.exception
    assert len(app.get("download_button")) == 2
    assert len(app.dataframe[0].value) == 2
    assert set(app.dataframe[0].value["Status"]) == {"Pago"}


def test_summary_ui_can_include_paid_pending_and_returns(monkeypatch):
    from src.data import cobranca_history, payment_history, devolucao_history
    base = records()
    monkeypatch.setattr(cobranca_history, "load_history", lambda: base.iloc[:2])
    monkeypatch.setattr(payment_history, "load_payments", lambda: base.iloc[2:])
    returned = base.iloc[2:].copy()
    returned["COD_LANCAMENTO"] = "C"
    monkeypatch.setattr(devolucao_history, "load_devolucoes", lambda: returned)
    app = AppTest.from_function(_summary_page, default_timeout=60)
    app.session_state["frame"] = base.rename(columns=HISTORY_LABELS)
    app.run()
    app.checkbox[0].check().run()
    assert not app.exception
    assert len(app.dataframe[0].value) == 3
    assert set(app.dataframe[0].value["Status"]) == {"Pendente", "Pago", "Devolução"}
